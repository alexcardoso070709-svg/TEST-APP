from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFFFCC")
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EUR = '#,##0.00" EUR"'
PCT = "0.0%"

wb = Workbook()

def style_header(ws, row, cols, text_list):
    for i, text in enumerate(text_list):
        c = ws.cell(row=row, column=cols[0] + i, value=text)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="left")

def label_input(ws, row, label, value, fmt=None, col_label=1, col_value=2):
    lc = ws.cell(row=row, column=col_label, value=label)
    lc.font = BLACK
    vc = ws.cell(row=row, column=col_value, value=value)
    vc.font = BLUE
    vc.fill = INPUT_FILL
    vc.border = BORDER
    if fmt:
        vc.number_format = fmt
    return f"{get_column_letter(col_value)}{row}"

def label_formula(ws, row, label, formula, fmt=None, col_label=1, col_value=2, bold=False):
    lc = ws.cell(row=row, column=col_label, value=label)
    lc.font = BOLD if bold else BLACK
    vc = ws.cell(row=row, column=col_value, value=formula)
    vc.font = BOLD if bold else BLACK
    vc.border = BORDER
    if fmt:
        vc.number_format = fmt
    return f"{get_column_letter(col_value)}{row}"

# ---------- Tab 0: Anleitung ----------
ws = wb.active
ws.title = "Anleitung"
ws.column_dimensions["A"].width = 90
ws["A1"] = "Freelancer-Finanzrechner-Set"
ws["A1"].font = TITLE
ws["A3"] = "Drei Rechner in diesem Workbook: Stundensatz-Kalkulator, Break-even-Rechner, Steuerruecklage-Rechner."
ws["A3"].font = BLACK
ws["A5"] = "Gelb hinterlegte Zellen = hier eigene Zahlen eintragen. Alle anderen Zellen rechnen automatisch."
ws["A5"].font = BOLD
ws["A7"] = "Beispielwerte sind bereits eingetragen (realistische Annahmen fuer einen Freelancer in Deutschland, 2026)."
ws["A7"].font = BLACK
ws["A9"] = "Hinweis: Steuersaetze und Sozialabgaben sind vereinfachte Annahmen, kein Steuerberater-Ersatz."
ws["A9"].font = Font(name=FONT, italic=True, color="666666")

# ---------- Tab 1: Stundensatz-Kalkulator ----------
ws = wb.create_sheet("Stundensatz-Kalkulator")
ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 18
ws["A1"] = "Stundensatz-Kalkulator"
ws["A1"].font = TITLE
style_header(ws, 3, (1,), ["Eingaben"])

r = 4
netto_row = label_input(ws, r, "Gewuenschtes Jahres-Nettoeinkommen", 42000, EUR); r += 1
ausgaben_row = label_input(ws, r, "Betriebsausgaben pro Jahr (Software, Buero, Versicherung)", 4800, EUR); r += 1
steuer_pct_row = label_input(ws, r, "Ruecklage Steuern (Anteil vom Nettoeinkommen)", 0.28, PCT); r += 1
vorsorge_pct_row = label_input(ws, r, "Ruecklage Kranken-/Rentenversicherung (Anteil vom Nettoeinkommen)", 0.20, PCT); r += 1
arbeitstage_row = label_input(ws, r, "Arbeitstage pro Jahr (ohne Wochenenden)", 260, "0"); r += 1
urlaub_row = label_input(ws, r, "Urlaubstage pro Jahr", 28, "0"); r += 1
krank_row = label_input(ws, r, "Erwartete Krankheits-/Ausfalltage pro Jahr", 10, "0"); r += 1
stunden_tag_row = label_input(ws, r, "Abrechenbare Stunden pro Arbeitstag", 5, "0.0"); r += 1

r += 1
style_header(ws, r, (1,), ["Ergebnis"]); r += 1
eff_tage_row = label_formula(ws, r, "Effektive Arbeitstage pro Jahr",
    f"={arbeitstage_row}-{urlaub_row}-{krank_row}", "0"); r += 1
stunden_jahr_row = label_formula(ws, r, "Abrechenbare Stunden pro Jahr",
    f"={eff_tage_row}*{stunden_tag_row}", "0"); r += 1
gesamtbedarf_row = label_formula(ws, r, "Gesamtbedarf pro Jahr (Netto + Ausgaben + Ruecklagen)",
    f"={netto_row}+{ausgaben_row}+({netto_row}*{steuer_pct_row})+({netto_row}*{vorsorge_pct_row})", EUR); r += 1
stundensatz_row = label_formula(ws, r, "Benoetigter Stundensatz",
    f"={gesamtbedarf_row}/{stunden_jahr_row}", EUR, bold=True); r += 1
tagessatz_row = label_formula(ws, r, "Entsprechender Tagessatz",
    f"={stundensatz_row}*{stunden_tag_row}", EUR, bold=True); r += 1

# ---------- Tab 2: Break-even-Rechner ----------
ws = wb.create_sheet("Break-even-Rechner")
ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 18
ws["A1"] = "Break-even-Rechner"
ws["A1"].font = TITLE
style_header(ws, 3, (1,), ["Eingaben"])

r = 4
fix_row = label_input(ws, r, "Fixkosten pro Monat (Miete, Software, Versicherung)", 650, EUR); r += 1
var_row = label_input(ws, r, "Variable Kosten pro Einheit/Kunde", 5, EUR); r += 1
preis_row = label_input(ws, r, "Verkaufspreis pro Einheit/Kunde", 25, EUR); r += 1

r += 1
style_header(ws, r, (1,), ["Ergebnis"]); r += 1
deckungsbeitrag_row = label_formula(ws, r, "Deckungsbeitrag pro Einheit (Preis - variable Kosten)",
    f"={preis_row}-{var_row}", EUR); r += 1
be_menge_row = label_formula(ws, r, "Break-even-Menge pro Monat (Einheiten)",
    f"={fix_row}/{deckungsbeitrag_row}", "0.0", bold=True); r += 1
be_umsatz_row = label_formula(ws, r, "Break-even-Umsatz pro Monat",
    f"={be_menge_row}*{preis_row}", EUR, bold=True); r += 1

r += 1
style_header(ws, r, (1,), ["Was-waere-wenn: bei X Einheiten/Monat"]); r += 1
einheiten_row = label_input(ws, r, "Angenommene verkaufte Einheiten pro Monat", 40, "0"); r += 1
umsatz_row = label_formula(ws, r, "Umsatz", f"={einheiten_row}*{preis_row}", EUR); r += 1
kosten_row = label_formula(ws, r, "Gesamtkosten", f"={fix_row}+({einheiten_row}*{var_row})", EUR); r += 1
gewinn_row = label_formula(ws, r, "Gewinn/Verlust", f"={umsatz_row}-{kosten_row}", EUR, bold=True); r += 1

# ---------- Tab 3: Steuerruecklage-Rechner ----------
ws = wb.create_sheet("Steuerruecklage-Rechner")
ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 18
ws["A1"] = "Steuerruecklage-Rechner"
ws["A1"].font = TITLE
style_header(ws, 3, (1,), ["Eingaben"])

r = 4
gewinn_jahr_row = label_input(ws, r, "Erwarteter Gewinn dieses Jahr", 35000, EUR); r += 1
steuersatz_row = label_input(ws, r, "Geschaetzter Steuersatz (Einkommensteuer + Soli, vereinfacht)", 0.27, PCT); r += 1
puffer_row = label_input(ws, r, "Sicherheitspuffer (falls Gewinn hoeher ausfaellt als erwartet)", 0.10, PCT); r += 1
bereits_gezahlt_row = label_input(ws, r, "Bereits geleistete Steuervorauszahlungen dieses Jahr", 0, EUR); r += 1

r += 1
style_header(ws, r, (1,), ["Ergebnis"]); r += 1
steuerlast_row = label_formula(ws, r, "Erwartete Steuerlast (mit Puffer)",
    f"={gewinn_jahr_row}*{steuersatz_row}*(1+{puffer_row})", EUR); r += 1
restbedarf_row = label_formula(ws, r, "Noch zurueckzulegender Betrag",
    f"={steuerlast_row}-{bereits_gezahlt_row}", EUR); r += 1
monatlich_row = label_formula(ws, r, "Monatlich zurueckzulegen (restliche 12 Monate)",
    f"={restbedarf_row}/12", EUR, bold=True); r += 1

wb.save("/home/user/TEST-APP/bundle/excel-rechner/Freelancer-Finanzrechner-Set.xlsx")
print("saved")
